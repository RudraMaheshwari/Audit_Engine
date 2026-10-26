"""
Generates AI_Auditor_Provisioning_Requirements.pptx
Run: python docs/generate_provisioning_slide.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree
import copy

# ── Slide dimensions (widescreen 16:9) ───────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

slide_layout = prs.slide_layouts[6]   # blank
slide = prs.slides.add_slide(slide_layout)

# ── Colours ───────────────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x1F, 0x49, 0x7D)
RED         = RGBColor(0xC0, 0x00, 0x00)
NAVY        = RGBColor(0x1F, 0x37, 0x64)
TEAL        = RGBColor(0x1F, 0x5C, 0x4A)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY  = RGBColor(0xF2, 0xF2, 0xF2)
MID_GREY    = RGBColor(0xD6, 0xDC, 0xE4)
GREEN_TXT   = RGBColor(0x37, 0x86, 0x10)
ORANGE_TXT  = RGBColor(0xC5, 0x5A, 0x11)
BLACK       = RGBColor(0x00, 0x00, 0x00)
GOLD        = RGBColor(0xFF, 0xC0, 0x00)
EXTRA_BG    = RGBColor(0xFF, 0xF2, 0xCC)   # yellow highlight for extra/different rows


def add_rect(slide, left, top, width, height, fill_rgb, text="", font_size=8,
             bold=False, font_color=WHITE, align=PP_ALIGN.CENTER,
             vertical_anchor=None, rotate=0):
    from pptx.util import Pt
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf    = txBox.text_frame
    tf.word_wrap = True
    if vertical_anchor:
        tf.vertical_anchor = vertical_anchor
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(font_size)
    run.font.bold  = bold
    run.font.color.rgb = font_color
    # fill
    sp   = txBox._element
    spPr = sp.find(qn("p:spPr"))
    solidFill = parse_xml(
        f'<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:srgbClr val="{fill_rgb[0]:02X}{fill_rgb[1]:02X}{fill_rgb[2]:02X}"/>'
        f'</a:solidFill>'
    )
    ln = parse_xml(
        '<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:noFill/></a:ln>'
    )
    spPr.append(solidFill)
    spPr.append(ln)
    if rotate:
        xfrm = spPr.find(qn("a:xfrm"))
        if xfrm is None:
            xfrm = parse_xml('<a:xfrm xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
            spPr.insert(0, xfrm)
        xfrm.set("rot", str(rotate * 60000))
    return txBox


def add_text(slide, left, top, width, height, text, font_size=8,
             bold=False, font_color=BLACK, align=PP_ALIGN.LEFT,
             bg_rgb=None, italic=False, wrap=True):
    from pptx.util import Pt
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf    = txBox.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = Pt(font_size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = font_color
    if bg_rgb:
        sp   = txBox._element
        spPr = sp.find(qn("p:spPr"))
        solidFill = parse_xml(
            f'<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:srgbClr val="{bg_rgb[0]:02X}{bg_rgb[1]:02X}{bg_rgb[2]:02X}"/>'
            f'</a:solidFill>'
        )
        ln = parse_xml(
            '<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:noFill/></a:ln>'
        )
        spPr.append(solidFill)
        spPr.append(ln)
    return txBox


# ── Layout constants ──────────────────────────────────────────────────────────
L  = Inches(0.15)    # left margin
T  = Inches(0.72)    # top of table (below title)
RH = Inches(0.30)    # row height
GW = Inches(0.28)    # group label column width

# Column widths
C0 = GW
C1 = Inches(1.55)   # Requirement
C2 = Inches(2.80)   # Usage
C3 = Inches(3.60)   # Tool
C4 = Inches(1.40)   # Availability
C5 = Inches(1.20)   # Access
# total ≈ 11.0" leaving margin

TOTAL_W = C0 + C1 + C2 + C3 + C4 + C5

# ── Title ─────────────────────────────────────────────────────────────────────
add_rect(slide, L, Inches(0.08), TOTAL_W, Inches(0.55),
         DARK_BLUE,
         "AI Auditor — TPI Due Diligence  |  Provisioning Requirements",
         font_size=14, bold=True)

# ── Legend note ───────────────────────────────────────────────────────────────
add_text(slide, L + TOTAL_W + Inches(0.08), Inches(0.08), Inches(2.2), Inches(0.55),
         "★ = Additional requirement beyond approved parallel-team stack",
         font_size=7, italic=True, font_color=RGBColor(0x7F, 0x60, 0x00))

# ── Header row ────────────────────────────────────────────────────────────────
hT   = T
cols = [("", C0), ("Requirement", C1), ("Usage", C2),
        ("Tool", C3), ("Availability", C4), ("Access", C5)]
x = L
for label, w in cols:
    add_rect(slide, x, hT, w, RH, MID_GREY, label,
             font_size=8, bold=True, font_color=DARK_BLUE, align=PP_ALIGN.CENTER)
    x += w

# ── Data rows ─────────────────────────────────────────────────────────────────
# Each row: (group, requirement, usage, tool, availability, access, extra)
# extra=True → yellow highlight, shows ★
ROWS = [
    # ── DATA PREPARATION ──────────────────────────────────────────────────────
    ("Data\nPrep.",
     "Document Reading,\nValidation & OCR",
     "Reading policy PDFs, supporting docs, scanned image files",
     "Python – pypdf2, python-docx, openpyxl\nTesseract OCR (scanned PDFs) ★",
     "Available with\nDatabricks",
     "Completed",
     False),
    ("Data\nPrep.",
     "Vector Search\nIndexes ★",
     "RAG retrieval for HCBI Policy, TPI DD Policy and Supporting Docs (3 indexes)",
     "Databricks Vector Search\n(Direct Access indexes — hcbi, policy, supporting_docs)",
     "Available with\nDatabricks",
     "Required",
     True),
    ("Data\nPrep.",
     "Database",
     "Questionnaire input tables, output Delta tables, scope embeddings, TPI inventory",
     "Databricks Unity Catalog / Delta Tables\n(cat_cbt_gaa_brz_dev)",
     "Available with\nDatabricks",
     "Completed",
     False),
    ("Data\nPrep.",
     "File Storage\n(Volumes)",
     "Store policy documents, supporting evidence, TPI inventory file, HCC GL map",
     "Databricks Volumes\n(/Volumes/cat_cbt_gaa_brz_dev/...)",
     "Available with\nDatabricks",
     "Completed",
     False),
    # ── MODEL DEVELOPMENT ─────────────────────────────────────────────────────
    ("Model\nDev.",
     "Large Language\nModels",
     "6-agent audit pipeline – risk classification, gap detection, policy analysis, NL-to-SQL",
     "Azure OpenAI – GPT-5 ★\n(JNJ Internal Gateway: genaiapimna.jnj.com)",
     "Required\n(JNJ Gateway)",
     "In Process",
     True),
    ("Model\nDev.",
     "Embeddings ★",
     "Document chunking, semantic vector search across policy and evidence indexes",
     "Azure OpenAI – text-embedding-3-small ★\n(Same JNJ gateway endpoint)",
     "Required\n(JNJ Gateway)",
     "In Process",
     True),
    ("Model\nDev.",
     "SQL Warehouse ★",
     "NL-to-SQL transactional queries against payment, credit memo and revenue tables",
     "Databricks SQL Warehouse ★\n(TransactionalDataAgent – live Delta table queries)",
     "Available with\nDatabricks",
     "Required",
     True),
    ("Model\nDev.",
     "Orchestration &\nSecrets",
     "Multi-agent parallel pipeline (asyncio, up to 5 concurrent), secrets management",
     "Python asyncio\nAzure Key Vault (Databricks secrets scope)",
     "Available /\nRequired",
     "Completed",
     False),
    ("Model\nDev.",
     "Source Code &\nVersion Tracking",
     "Source code version control, branching and merging",
     "Bitbucket",
     "Required",
     "Completed",
     False),
    # ── DEPLOYMENT ────────────────────────────────────────────────────────────
    ("Deploy.",
     "Deployment\nEnvironments / CI/CD",
     "Azure Cloud environments for DEV, UAT and PROD",
     "Azure DEV, UAT, PROD\nJenkins for CI/CD",
     "Required",
     "Completed",
     False),
    ("Deploy.",
     "Docker & Container\nScanning",
     "Securely stores and manages Docker images and scans for vulnerabilities",
     "Jrog\nFortify / CheckMarx",
     "Required",
     "Completed",
     False),
    ("Deploy.",
     "Application Hosting\n& Load Balancing",
     "Frontend dashboard (Streamlit) + backend pipeline hosting and load management",
     "Load Balancer Controller\nStreamlit (frontend UI)",
     "Required",
     "Completed",
     False),
    ("Deploy.",
     "Infra as Code &\nMonitoring",
     "Monitor pipeline runs, audit job tracking, infra metrics and consumption",
     "Application Insights, Azure Monitor\nDatabricks Job Monitoring, Terraform",
     "Required",
     "Completed",
     False),
]

# Group boundaries and colors
GROUP_RANGES = {
    "Data\nPrep.":  (0, 3,  RED),
    "Model\nDev.":  (4, 8,  NAVY),
    "Deploy.":      (9, 12, TEAL),
}

# Draw rows
row_tops = []
y = T + RH
for i, row in enumerate(ROWS):
    row_tops.append(y)
    y += RH

# Group side bars + labels
for grp_label, (r_start, r_end, color) in GROUP_RANGES.items():
    grp_top = row_tops[r_start]
    grp_bot = row_tops[r_end] + RH
    grp_h   = grp_bot - grp_top
    add_rect(slide, L, grp_top, GW, grp_h, color,
             grp_label, font_size=7, bold=True, font_color=WHITE, align=PP_ALIGN.CENTER)

# Row cells
for i, (grp, req, usage, tool, avail, access, extra) in enumerate(ROWS):
    y    = row_tops[i]
    bg   = EXTRA_BG if extra else (LIGHT_GREY if i % 2 == 0 else WHITE)
    x    = L + C0

    # Requirement
    txt = ("★ " if extra else "") + req
    add_text(slide, x, y, C1, RH, txt, font_size=7.5, bold=True,
             font_color=DARK_BLUE if not extra else RGBColor(0x7F, 0x60, 0x00),
             bg_rgb=bg, align=PP_ALIGN.LEFT)
    x += C1

    # Usage
    add_text(slide, x, y, C2, RH, usage, font_size=7, bg_rgb=bg)
    x += C2

    # Tool
    add_text(slide, x, y, C3, RH, tool, font_size=7,
             font_color=NAVY, bg_rgb=bg, bold=False)
    x += C3

    # Availability
    add_text(slide, x, y, C4, RH, avail, font_size=7, bg_rgb=bg,
             align=PP_ALIGN.CENTER, font_color=BLACK)
    x += C4

    # Access
    a_color = GREEN_TXT if access == "Completed" else (ORANGE_TXT if "Process" in access else RED)
    add_text(slide, x, y, C5, RH, access, font_size=7.5, bold=True,
             font_color=a_color, bg_rgb=bg, align=PP_ALIGN.CENTER)

# ── Footer note ───────────────────────────────────────────────────────────────
footer_y = row_tops[-1] + RH + Inches(0.05)
add_text(slide, L, footer_y, TOTAL_W, Inches(0.28),
         "★ Items marked are additional to the approved parallel-team stack and require separate provisioning sign-off.  "
         "All other items align with already-approved services.",
         font_size=6.5, italic=True, font_color=RGBColor(0x60, 0x60, 0x60),
         bg_rgb=RGBColor(0xFF, 0xFF, 0xFF))

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"F:\EXL\AI_Auditor\docs\AI_Auditor_Provisioning_Requirements.pptx"
prs.save(out)
print(f"Saved: {out}")
