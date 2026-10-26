"""Per-scope proposed sample size, read from the CPR guidance workbook.

The guidance workbook (CPR AI - System Guide.xlsx) has one sheet per scope area, with the
expected sample size in column K (row 4 onward). We index that number by both the scope_id and the
normalised sheet name so a lookup succeeds whether Output C stores a scope_id or a display name.

No buffer is added — the number is returned exactly as the guidance states it.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import openpyxl

from src.config.settings import Settings
from src.ingestion.extractors.cpr_scope_extractor import _resolve_scope_id
from src.utils.logger import get_logger

from .scope_key import normalize

logger = get_logger(__name__)


def _looks_like_volume(path: str) -> bool:
    """True for a Databricks Unity Catalog Volumes path (only readable via the Databricks API)."""
    p = str(path).replace("\\", "/").lower()
    return "/volumes/" in p or p.startswith("volumes/")


def _open_workbook(path: str, settings: Settings | None):
    """Return a loaded openpyxl workbook, fetching from a Databricks Volume when needed."""
    if _looks_like_volume(path) and settings is not None:
        # Normalise Windows backslashes back to a proper /Volumes/... path for the Files API.
        volume_path = "/" + str(path).replace("\\", "/").lstrip("/")
        from databricks.sdk import WorkspaceClient

        from src.config.databricks_auth import build_config

        client = WorkspaceClient(config=build_config(settings))
        resp = client.files.download(volume_path)
        data = resp.contents.read()
        logger.info("[Guidance] Downloaded guidance workbook from volume %s", volume_path)
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)

    return openpyxl.load_workbook(Path(path), read_only=True, data_only=True)

_COL_EXPECTED_SAMPLE_SIZE = 11  # column K
_HEADER_ROWS = 3                # rows 1-3 are group/column/description headers; data starts at row 4


def _first_int(text: str) -> int | None:
    m = re.search(r"\d+", text or "")
    return int(m.group()) if m else None


class GuidanceSampleSizes:
    """Lookup of scope area -> proposed sample size from the CPR guidance workbook."""

    def __init__(self) -> None:
        self._by_scope_id: dict[str, int] = {}
        self._by_norm_name: dict[str, int] = {}

    @classmethod
    def load(cls, path: str | Path, settings: Settings | None = None) -> "GuidanceSampleSizes":
        inst = cls()
        try:
            wb = _open_workbook(str(path), settings)
        except Exception as exc:
            logger.warning(
                "[Guidance] Could not open guidance workbook %s — sample sizes will be blank | %s",
                path, exc,
            )
            return inst

        try:
            for sheet_name in wb.sheetnames:
                if sheet_name.strip().lower() == "index":
                    continue
                ws = wb[sheet_name]
                sample = inst._read_sample_size(ws)
                if sample is None:
                    continue
                inst._by_norm_name[normalize(sheet_name)] = sample
                scope_id = _resolve_scope_id(sheet_name)
                if scope_id:
                    inst._by_scope_id[scope_id] = sample
        finally:
            wb.close()

        logger.info("[Guidance] Loaded sample sizes for %d scope sheet(s)", len(inst._by_norm_name))
        return inst

    def _read_sample_size(self, ws) -> int | None:
        for row in ws.iter_rows(
            min_row=_HEADER_ROWS + 1,
            min_col=_COL_EXPECTED_SAMPLE_SIZE,
            max_col=_COL_EXPECTED_SAMPLE_SIZE,
            values_only=True,
        ):
            val = row[0]
            if val is None:
                continue
            n = _first_int(str(val))
            if n is not None:
                return n
        return None

    def lookup(self, scope_area: str) -> int | None:
        """Return the proposed sample size for a scope area, or None if not found."""
        # Direct scope_id (Output C stored an id like "TENDERS").
        if scope_area in self._by_scope_id:
            return self._by_scope_id[scope_area]
        # Resolve display name / id via the shared scope resolver.
        scope_id = _resolve_scope_id(scope_area)
        if scope_id and scope_id in self._by_scope_id:
            return self._by_scope_id[scope_id]
        # Normalised display-name match.
        return self._by_norm_name.get(normalize(scope_area))
