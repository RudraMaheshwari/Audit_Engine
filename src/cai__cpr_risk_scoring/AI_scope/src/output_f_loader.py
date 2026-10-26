"""Load the per-scope transaction totals (Output F — Scope Transaction Summary) for AI Scope's story.

Gives, per scope area: Total Transactions + Total Amount. AI Scope uses these for the FIRST line of a
scope's story and to compute its share of total transactional spend. BEST-EFFORT: if the source is
absent/unreadable the loader returns ({}, 0.0) and the story simply omits line 1 for that scope.

Source is DERIVED from the Output C source (same catalog.schema / same workbook) + engagement — it is
the existing risk-analysis Output F, not a new input the caller must pass.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from sqlalchemy import text

from src.utils.logger import get_logger

from .config import SOURCE_EXCEL, AIScopeConfig
from .scope_key import normalize

logger = get_logger(__name__)


@dataclass
class ScopeTxn:
    total_txns: int | None
    amount_usd: float | None
    amount_display: str   # the original "Local / USD" amount string, for display


# normalised header -> field
_TXN_FIELDS = {
    "scopearea": "scope_area",
    "totaltransactions": "total_txns",
    "totalamountlocalusd": "total_amount",
    "totalamount": "total_amount",
}


def _parse_int(value) -> int | None:
    m = re.search(r"\d[\d,]*", str(value or ""))
    return int(m.group().replace(",", "")) if m else None


def _parse_usd(value) -> float | None:
    """Pull the USD figure out of a 'AUD 1,234.00 / USD 800.00' style amount; fall back to the last number."""
    s = str(value or "")
    m = re.search(r"USD\s*([\d,]+(?:\.\d+)?)", s)
    if m:
        return float(m.group(1).replace(",", ""))
    nums = re.findall(r"[\d,]+(?:\.\d+)?", s)
    return float(nums[-1].replace(",", "")) if nums else None


def _derive_db_table(config: AIScopeConfig) -> str:
    """Sibling Output F table in the same catalog.schema as the Output C table."""
    prefix = config.path.rsplit(".", 1)[0]  # catalog.schema
    return f"{prefix}.Scope_Transaction_Summary_{config.settings.engagement_slug}"


def _records_from_db(config: AIScopeConfig):
    from src.config.databricks_auth import build_engine

    # Explicit Output F table (AI_SCOPE_OUTPUT_F_PATH) wins; otherwise derive the sibling table name.
    table = config.output_f_path or _derive_db_table(config)
    engine = build_engine(config.settings)
    with engine.connect() as conn:
        result = conn.execute(text(f"SELECT * FROM {table}"))
        cols = [str(c) for c in result.keys()]
        rows = result.fetchall()
    field_by_idx = {i: _TXN_FIELDS.get(normalize(c)) for i, c in enumerate(cols)}
    field_by_idx = {i: f for i, f in field_by_idx.items() if f}
    return [{f: r[i] for i, f in field_by_idx.items()} for r in rows]


def _records_from_excel(config: AIScopeConfig):
    """Read the scope-transaction-summary sheet from the SAME workbook as Output C."""
    wb = openpyxl.load_workbook(Path(config.path), read_only=True, data_only=True)
    try:
        ws = None
        for name in wb.sheetnames:
            nl = name.lower()
            if "txn" in nl or "transaction summary" in nl or "output f" in nl:
                ws = wb[name]
                break
        if ws is None:
            return []
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    field_by_idx = {i: _TXN_FIELDS.get(normalize(str(h or ""))) for i, h in enumerate(rows[0])}
    field_by_idx = {i: f for i, f in field_by_idx.items() if f}
    return [
        {f: (r[i] if i < len(r) else None) for i, f in field_by_idx.items()}
        for r in rows[1:]
    ]


def load_scope_txn_totals(config: AIScopeConfig) -> tuple[dict[str, ScopeTxn], float]:
    """Return ({normalised scope -> ScopeTxn}, grand_total_usd). Best-effort — ({}, 0.0) on any failure."""
    try:
        records = _records_from_excel(config) if config.source == SOURCE_EXCEL else _records_from_db(config)
    except Exception as exc:  # noqa: BLE001 — story just omits line 1 if totals are unavailable
        logger.warning("[OutputF] Scope transaction totals unavailable — line 1 skipped | %s", exc)
        return {}, 0.0

    out: dict[str, ScopeTxn] = {}
    grand = 0.0
    for rec in records:
        scope = str(rec.get("scope_area") or "").strip()
        if not scope:
            continue
        usd = _parse_usd(rec.get("total_amount"))
        out[normalize(scope)] = ScopeTxn(
            total_txns=_parse_int(rec.get("total_txns")),
            amount_usd=usd,
            amount_display=str(rec.get("total_amount") or "").strip(),
        )
        if usd:
            grand += usd
    logger.info("[OutputF] Loaded transaction totals for %d scope(s) | grand_total_usd=%.0f", len(out), grand)
    return out, grand
