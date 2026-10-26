"""Load Output E (Past Audit Report) observations from the generated workbook.

Direction-2's "Past Audit" source uses these curated, summarised observations instead of re-scanning
the raw past_audit vector index per scope — it is far faster and already filtered to the risky scopes.
Only observations from the last N years (default 3) are kept, so stale findings never drive a risk.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy import text

from src.config.constants import SHEET_E_NAME
from src.utils.logger import get_logger

from .config import SOURCE_EXCEL, AIScopeConfig
from .scope_key import normalize

logger = get_logger(__name__)

_RECENCY_YEARS = 3

# Normalised header -> internal field (so display headers resolve regardless of exact spelling).
_FIELD_BY_KEY = {
    "scopearea": "scope_area",
    "issuetype": "issue_type",
    "issuedescription": "issue_description",
    "reportdate": "report_date",
    "criticality": "criticality",
    "auditreport": "audit_report",
    "sourcedocument": "audit_report",
}


@dataclass
class PastReviewObs:
    scope_area: str
    issue_type: str
    issue_description: str
    report_date: str
    audit_report: str = ""


def _within_recency(report_date: str, min_year: int) -> bool:
    """True when the report date's year is >= min_year. Undated rows are kept (can't prove stale)."""
    m = re.search(r"(19|20)\d{2}", str(report_date or ""))
    if not m:
        return True
    return int(m.group()) >= min_year


def _records_from_excel(config: AIScopeConfig) -> list[dict]:
    """Read the Output E (Past Audit) sheet from the same workbook as Output C → list of field dicts."""
    path = Path(config.path)
    if not path.exists():
        logger.warning("[OutputE] Workbook not found: %s", path)
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        logger.warning("[OutputE] Could not open workbook %s | %s", path, exc)
        return []
    try:
        ws = next((wb[s] for s in wb.sheetnames if normalize(s) == normalize(SHEET_E_NAME)), None)
        if ws is None:
            logger.info("[OutputE] No '%s' sheet in workbook — skipping Past Audit source", SHEET_E_NAME)
            return []
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    col_field = {
        idx: _FIELD_BY_KEY[normalize(str(h or ""))]
        for idx, h in enumerate(rows[0])
        if normalize(str(h or "")) in _FIELD_BY_KEY
    }
    if "scope_area" not in col_field.values():
        logger.warning("[OutputE] Output E sheet missing a Scope Area column — headers=%s", list(rows[0]))
        return []
    return [
        {field: (row[idx] if idx < len(row) else None) for idx, field in col_field.items()}
        for row in rows[1:]
    ]


def _records_from_db(config: AIScopeConfig) -> list[dict]:
    """Read the Output E (Past Audit) Delta table named by AI_SCOPE_OUTPUT_E_PATH → list of field dicts."""
    from src.config.databricks_auth import build_engine

    table = config.output_e_path
    engine = build_engine(config.settings)
    with engine.connect() as conn:
        result = conn.execute(text(f"SELECT * FROM {table}"))
        cols = [str(c) for c in result.keys()]
        rows = result.fetchall()
    col_field = {i: _FIELD_BY_KEY.get(normalize(c)) for i, c in enumerate(cols)}
    col_field = {i: f for i, f in col_field.items() if f}
    if "scope_area" not in col_field.values():
        logger.warning("[OutputE] Output E table %r missing a Scope Area column — columns=%s", table, cols)
        return []
    return [{f: r[i] for i, f in col_field.items()} for r in rows]


def load_output_e(config: AIScopeConfig) -> list[PastReviewObs]:
    """Return last-3-years Past Audit observations from Output E (workbook sheet for the excel source, or
    the AI_SCOPE_OUTPUT_E_PATH Delta table for the db_table source). [] on failure / when unavailable."""
    try:
        if config.source == SOURCE_EXCEL:
            records = _records_from_excel(config)
        elif config.output_e_path:
            records = _records_from_db(config)
        else:
            logger.info("[OutputE] db_table source without AI_SCOPE_OUTPUT_E_PATH — Past Audit skipped")
            return []
    except Exception as exc:  # noqa: BLE001 — past-audit source is best-effort
        logger.warning("[OutputE] Past Audit observations unavailable | %s", exc)
        return []

    min_year = datetime.now().year - _RECENCY_YEARS
    out: list[PastReviewObs] = []
    skipped_old = 0
    for rec in records:
        scope = str(rec.get("scope_area") or "").strip()
        issue_type = str(rec.get("issue_type") or "").strip()
        if not scope or not issue_type:
            continue
        report_date = str(rec.get("report_date") or "").strip()
        if not _within_recency(report_date, min_year):
            skipped_old += 1
            continue
        out.append(PastReviewObs(
            scope_area=scope,
            issue_type=issue_type,
            issue_description=str(rec.get("issue_description") or "").strip(),
            report_date=report_date,
            audit_report=str(rec.get("audit_report") or "").strip(),
        ))
    logger.info(
        "[OutputE] Loaded %d Past Audit observation(s) within last %d years (dropped %d older)",
        len(out), _RECENCY_YEARS, skipped_old,
    )
    return out
