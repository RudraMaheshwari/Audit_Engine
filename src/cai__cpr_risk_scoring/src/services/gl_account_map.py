"""Loader for the HCC GL Account Map.

The workbook maps GL account codes to their approved HCC activity / description. The transactional
analysis agent uses it to verify whether a GL code used on a transaction is the approved account for
the activity under review, rather than asserting miscoding from the raw data alone.

Reads the GL-code column and the description/activity column from the configured map file and returns
the (code, description) pairs. Degrades gracefully to an empty mapping when the file, sheet, or columns
are missing — the analysis prompt then treats mapping as unconfirmed rather than failing.
"""

from __future__ import annotations

import openpyxl

from src.config.settings import Settings
from src.utils.logger import get_logger

logger = get_logger(__name__)

_CODE_HEADER_FRAGMENTS = ("gl account code", "gl account", "account code", "gl code", "gl", "account")
_DESC_HEADER_FRAGMENTS = ("description", "activity", "hcc area", "hcc", "area", "category", "name")

_MAX_ROWS = 2000


def _header_cells(header: tuple) -> list[str]:
    return [("" if h is None else str(h)).strip() for h in header]


def _find_column(header: tuple, fragments: tuple[str, ...], exclude: int | None = None) -> int | None:
    cells = [c.lower() for c in _header_cells(header)]
    for frag in fragments:
        for i, c in enumerate(cells):
            if i == exclude:
                continue
            if frag in c:
                return i
    return None


def load_gl_account_map(settings: Settings) -> list[tuple[str, str]]:
    """Return de-duplicated (GL code, description) pairs from the GL account map file.

    Column detection is schema-agnostic: the GL-code column is matched on header fragments, the
    description column on its own fragments (never the same column as the code). Returns an empty
    list (and logs a warning) if the file, sheet, or code column is missing.
    """
    path = settings.hcc_gl_account_map_path
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        logger.warning("[GLAccountMap] Could not open %s: %s", path, exc)
        return []
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            logger.warning("[GLAccountMap] Map file %s has no header row", path)
            return []

        code_idx = _find_column(header, _CODE_HEADER_FRAGMENTS)
        if code_idx is None:
            logger.warning("[GLAccountMap] GL code column not found | header=%s", _header_cells(header))
            return []
        desc_idx = _find_column(header, _DESC_HEADER_FRAGMENTS, exclude=code_idx)

        mapping: list[tuple[str, str]] = []
        seen: set[str] = set()
        for r in rows:
            if code_idx >= len(r):
                continue
            code_val = r[code_idx]
            if code_val is None:
                continue
            code = str(code_val).strip()
            if not code or code.lower() in seen:
                continue
            desc = ""
            if desc_idx is not None and desc_idx < len(r) and r[desc_idx] is not None:
                desc = str(r[desc_idx]).strip()
            seen.add(code.lower())
            mapping.append((code, desc))
            if len(mapping) >= _MAX_ROWS:
                break
        logger.info("[GLAccountMap] Loaded %d GL account mappings | file=%s", len(mapping), path)
        return mapping
    finally:
        wb.close()


def format_gl_account_map_block(mapping: list[tuple[str, str]]) -> str:
    """Format the GL account map as a context block for the transactional analysis prompt."""
    if not mapping:
        return ""
    lines = "\n".join(f"- {code}: {desc}" if desc else f"- {code}" for code, desc in mapping)
    return (
        "=== HCC GL ACCOUNT MAP — APPROVED GL ACCOUNTS PER ACTIVITY ===\n"
        "Verify a GL account used on a transaction against this authoritative map before concluding "
        "miscoding. A GL code present here is an APPROVED account for its listed activity. Conclude "
        "miscoding ONLY when the activity under review uses a GL code that this map assigns to a "
        "different activity. If the relevant activity or GL code is NOT in this map, the mapping is "
        "UNCONFIRMED — describe the issue as a coding inconsistency / potential miscoding, not "
        "confirmed incorrect coding.\n"
        f"{lines}"
    )
