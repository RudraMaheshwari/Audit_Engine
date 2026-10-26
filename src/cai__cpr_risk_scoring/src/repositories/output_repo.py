import logging
import os
import time
from pathlib import Path

import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.utils.dateformat import format_date, normalize_dates
from src.utils.xml_clean import sanitize_xlsx, xml_safe

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    _RICH_TEXT_OK = True
except ImportError:  # openpyxl < 3.1 — fall back to plain text
    _RICH_TEXT_OK = False

_MARKDOWN_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

def _xml_safe(value):
    """Strip illegal XML chars AND truncate to Excel's 32,767-char cell limit (over-length strings
    are the other cause of Excel's 'repaired String properties' prompt). Non-strings pass through."""
    return xml_safe(value)

_AMOUNT_RE = re.compile(
    r"((?:USD|AUD|NZD|EUR|GBP|US\$)\s?\$?[\d,]+(?:\.\d+)?(?:\s?(?:million|billion|thousand|bn))?"
    r"|\$\s?[\d,]+(?:\.\d+)?(?:\s?(?:million|billion|thousand|bn))?"
    r"|\b\d{1,3}(?:,\d{3})+(?:\.\d+)?\b"
    r"|\b\d+(?:\.\d+)?\s?(?:million|billion|thousand|bn)\b"
    r"|\b\d+(?:\.\d+)?%"
    r"|\b\d+(?:\.\d+)?x\b)",
    re.IGNORECASE,
)


def _apply_bold(text):
    """Return CellRichText with bold applied.

    Primary: bold anything wrapped in **...**  (LLM-generated markers).
    Fallback: if no markdown markers found, bold monetary/numeric amounts via regex.
    Returns plain string when rich-text is unsupported or no bold targets exist.
    """
    text = _xml_safe(text)
    if not text or not _RICH_TEXT_OK:
        return text

    use_re = _MARKDOWN_BOLD_RE if _MARKDOWN_BOLD_RE.search(text) else _AMOUNT_RE

    # Build the runs, SKIPPING any empty run. An empty plain segment or an empty bold (e.g. the LLM
    # emitting "****") writes an empty <r><t></t></r>, which makes Excel "repair String properties"
    # and strip the rich text (collapsing bulleted Description/Rationale into one paragraph).
    parts: list = []
    last = 0
    for m in use_re.finditer(text):
        if m.start() > last:
            seg = text[last:m.start()]
            if seg:
                parts.append(seg)
        bold_text = m.group(1) if use_re is _MARKDOWN_BOLD_RE else m.group(0)
        if bold_text and bold_text.strip():
            parts.append(TextBlock(InlineFont(b=True), bold_text))
        elif bold_text:  # whitespace-only — keep it as a plain run so spacing survives
            parts.append(bold_text)
        last = m.end()
    if last < len(text):
        tail = text[last:]
        if tail:
            parts.append(tail)
    if not parts or all(not isinstance(p, TextBlock) for p in parts):
        return text  # nothing actually bolded → plain string (no rich-text XML at all)
    return CellRichText(*parts)


def _format_source_details(text: str) -> str:
    """Put every 'Source:' entry on its OWN line as a plain string with real newlines.

    Splits on every 'Source:' boundary (works even when the text arrives as one paragraph joined
    by ';' / '|' / run-on bullets), strips stray bullet/pipe/semicolon residue, and rejoins with
    real '\\n' newlines so Excel renders each entry on its own line (with wrap_text). Plain string
    is used deliberately — openpyxl rich-text newlines do not reliably render in Excel.
    """
    if not text or "Source:" not in text:
        return text
    entries = []
    for part in re.split(r"(?=Source:)", text):
        part = re.sub(r"^[\s•|;]+", "", part)
        part = re.sub(r"[\s•|;]+$", "", part)
        if not part.startswith("Source:"):
            continue
        # The questionnaire itself is not a citable policy/evidence source — drop the entry.
        if re.match(r"Source:\s*Questionnaire\b", part, flags=re.IGNORECASE):
            continue
        if re.search(r"File Name\s*:\s*Not Available", part, flags=re.IGNORECASE):
            continue
        # Drop 'Not Available' Section/Page fragments (e.g. spreadsheet policies with no page/section).
        part = re.sub(
            r"\s*[;|]\s*(?:Section|Page No\.?)\s*:\s*Not Available\.?",
            "",
            part,
            flags=re.IGNORECASE,
        )
        part = part.rstrip(";").strip()
        if part.startswith("Source:"):
            entries.append("• " + part)
    return "\n".join(entries) if entries else text

from src.config.constants import (
    CATEGORY_RANK,
    EXCEL_MAX_ROW_HEIGHT,
    SHEET_A_NAME,
    SHEET_B_NAME,
    SHEET_C_NAME,
    SHEET_D_NAME,
    SHEET_E_NAME,
    SHEET_F_NAME,
)
from src.models.schemas import OutputRow, RiskLevel, ScopeAreaSummary

logger = logging.getLogger(__name__)

# Row fill is driven by the Risk Category column:
#   High Risk           → red
#   Medium / Low Risk   → yellow
#   No Risk / unknown   → no fill (None)
_CATEGORY_RED    = "FFC7CE"
_CATEGORY_YELLOW = "FFEB9C"


def _category_fill(category: str) -> str | None:
    c = (category or "").strip().lower()
    if "high" in c:
        return _CATEGORY_RED
    if "medium" in c or "low" in c:
        return _CATEGORY_YELLOW
    return None  # No Risk / unknown → leave cell unfilled

_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=9)

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_OUTPUT_A_COLUMNS: list[tuple[str, str, int, bool]] = [
    ("sno",                    "#",                      5,  False),
    ("source_file",            "Questionnaire File",     28, True),
    ("qa_block",               "Questions & Responses",  80, True),
    ("processed_at",           "Processed At",           20, False),
    ("source_details",         "Source Details",         55, True),
    ("scope_area",             "Scope Area",             20, True),
    ("scope_area_reasoning",   "Scope Area Reasoning",   45, True),
    ("risk_type",              "Risk Type",              34, True),
    ("risk_rationale",         "Risk Rationale",         60, True),
    ("risk_description",       "Risk Description",       60, True),
    ("risk_category",          "Risk Category",          18, False),
    ("risk_category_reasoning","Risk Category Reasoning",50, True),
    ("risk_title",             "Risk Title",             22, True),
    ("evidence_explanation",   "Evidence Explanation",   45, True),
    ("calculation_evidence",   "Calculation Evidence",   45, True),
    ("sql_queries",            "SQL Queries",            50, True),
    ("samples",                "Samples",                60, True),
]


def _qa_block(row: OutputRow) -> str:
    """Compose the merged Output A 'Questions & Responses' cell: a '[<#>] <Questionnaire File>' header
    line (like Output C) on top, then Q / Response / Follow-up Q / Follow-up Response. Blank parts are
    omitted."""
    header = f"[{getattr(row, 'sno', '')}] {(getattr(row, 'source_file', '') or '').strip()}".strip()
    parts: list[str] = [header] if header.strip("[] ") else []
    for label, attr in (
        ("Q", "question"),
        ("Response", "response"),
        ("Follow-up Q", "followup_question"),
        ("Follow-up Response", "followup_response"),
    ):
        val = (getattr(row, attr, "") or "").strip()
        if val:
            parts.append(f"{label}: {val}")
    return "\n".join(parts)


def _rationale_with_category(row: OutputRow) -> str:
    """Output A 'Risk Rationale' = the risk_rationale bullets PLUS one final bullet carrying the
    Risk Category Reasoning. This augments only the Output A cell (and its Delta export, which reuses
    this helper); the stored risk_rationale and Output C are unchanged. The reasoning is added as the
    LAST bullet, with the same '• ' prefix as the other bullets, and is not duplicated if it is already
    present in the rationale."""
    rationale = (getattr(row, "risk_rationale", "") or "").strip()
    reasoning = (getattr(row, "risk_category_reasoning", "") or "").strip()
    if not reasoning or reasoning in rationale:
        return rationale
    bullet = reasoning if reasoning.lstrip().startswith(("•", "-", "*")) else f"• {reasoning}"
    return f"{rationale}\n{bullet}" if rationale else bullet


# Columns whose text must NOT be date-normalised (raw query text — rewriting a date inside SQL would
# misrepresent the query).
_NO_DATE_NORMALIZE = {"sql_queries"}


def _stack_bilingual(english: str, local: str | None) -> str:
    """Output A bilingual cell: the English answer, then the local-language rendering below it. Returns
    the English unchanged when no local translation is present (the normal English-only case)."""
    local = (local or "").strip()
    if not local:
        return english
    return f"{english}\n\n──────────\n{local}"


def _cell_value(row: OutputRow, attr: str):
    # Output A merges Q&A into ONE column. This synthetic attr is resolved here so BOTH the Output A
    # sheet and the Output A Delta export (which reuse this helper) stay consistent. Output C is
    # unaffected — it composes its own `questions_and_responses` field during aggregation.
    # Every string result is date-normalised to "Month D, YYYY" (except the SQL column); datetime
    # cells are formatted the same way.
    if attr == "qa_block":
        return normalize_dates(_qa_block(row))
    # Output A only: append the Risk Category Reasoning as the last bullet of the Risk Rationale cell,
    # then stack the local-language rendering (bilingual) when one was produced.
    if attr == "risk_rationale":
        english = normalize_dates(_rationale_with_category(row))
        return _stack_bilingual(english, getattr(row, "risk_rationale_local", None))
    raw = getattr(row, attr, None)
    if raw is None:
        return ""
    if hasattr(raw, "strftime"):
        return format_date(raw)
    if isinstance(raw, float):
        return round(raw, 2)
    text = str(raw)
    text = text if attr in _NO_DATE_NORMALIZE else normalize_dates(text)
    # Output A only: Risk Description is rendered bilingually when a local translation exists.
    if attr == "risk_description":
        return _stack_bilingual(text, getattr(row, "risk_description_local", None))
    return text


# Third sheet — one row per aggregated (clustered) risk. attr, header, width.
_OUTPUT_C_COLUMNS: list[tuple[str, str, int]] = [
    ("sno",                     "#",                      5),
    ("scope_area",              "Scope Area",             24),
    ("risk_type",               "Risk Type",              34),
    ("risk_category",           "Risk Category",          16),
    ("overall_risk_category",   "Overall Risk Category",  18),
    ("likelihood",              "Likelihood",             10),
    ("impact",                  "Impact",                 10),
    ("risk_description",        "Risk Description",       60),
    ("risk_rationale",          "Risk Rationale",         55),
    ("calculation_evidence",    "Calculation Evidence",   45),
    ("sql_queries",             "SQL Queries",            50),
    ("past_occurrence",         "Past Occurrence",        18),
    ("questionnaire_sources",   "Questionnaire Sources",  28),
    ("source_details",          "Source Details",         50),
    ("questions_and_responses", "Questions & Responses",  60),
]

_RISK_SUMMARY_COLUMNS: list[tuple[str, int]] = [
    ("Scope Area",            24),
    ("Overall Risk Category", 20),
    ("Risk Type",             34),
    ("Risk Category",         16),
    ("Past Occurrence",       18),
    ("Source",                20),
]

# Past Audit Report sheet — one row per observation. attr, header, width.
_OUTPUT_E_COLUMNS: list[tuple[str, str, int]] = [
    ("sno",               "#",                 5),
    ("source_document",   "Source Document",   32),
    ("report_date",       "Report Date",       16),
    ("scope_area",        "Scope Area",        24),
    ("issue_type",        "Issue Type",        28),
    ("issue_description", "Issue Description", 70),
    ("criticality",       "Criticality",       14),
    ("audit_report",      "Audit Report",      32),
]


def _agg_cell_value(risk, attr: str):
    if attr == "past_occurrence":
        if getattr(risk, "past_occurrence", False):
            link = getattr(risk, "past_occurrence_link", "") or ""
            return f"Yes — {link}" if link else "Yes"
        if not getattr(risk, "past_occurrence_determinable", True):
            return "Unable to determine based on available evidence"
        return "No"
    raw = getattr(risk, attr, None)
    if raw is None:
        return ""
    if hasattr(raw, "strftime"):
        return format_date(raw)
    text = str(raw)
    return text if attr in _NO_DATE_NORMALIZE else normalize_dates(text)


class OutputRepository:
    def __init__(self) -> None:
        self._wb = None
        self._ws_a = None
        self._active_path: Path | None = None

    def _ensure_workbook(self, path: Path):
        if self._wb is not None and self._active_path == path:
            return
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists():
            self._wb = openpyxl.load_workbook(path)
            self._ws_a = self._wb[SHEET_A_NAME]
        else:
            self._wb = openpyxl.Workbook()
            self._ws_a = self._wb.active
            self._ws_a.title = SHEET_A_NAME
            self._write_headers(self._ws_a)
            self._ws_a.freeze_panes = "A2"
            for col_idx, (_, _, width, _wrap) in enumerate(_OUTPUT_A_COLUMNS, start=1):
                self._ws_a.column_dimensions[
                    openpyxl.utils.get_column_letter(col_idx)
                ].width = width
        self._active_path = path

    def _atomic_save(self, path: Path) -> None:
        tmp = path.with_suffix(path.suffix + ".tmp")
        try:
            self._wb.save(tmp)
            sanitize_xlsx(tmp)  # permanent guarantee: clean the final bytes before publishing
        except Exception as exc:
            logger.warning("Output temp-save failed — will retry next row | %s", exc)
            return
        for _attempt in range(5):
            try:
                os.replace(tmp, path)
                return
            except PermissionError:
                time.sleep(0.4)
        logger.warning(
            "Output file is locked (is it open in Excel?) — skipping this save; the "
            "row is kept in memory and will be written on the next save | %s", path.name,
        )
        try:
            os.remove(tmp)
        except OSError:
            pass

    def append_row(self, row: OutputRow, output_path: str | Path) -> None:
        path = Path(output_path)
        self._ensure_workbook(path)
        self._write_data_row(self._ws_a, row, self._ws_a.max_row + 1)
        self._atomic_save(path)

    def write_summary(self, summary: ScopeAreaSummary, output_path: str | Path) -> None:
        path = Path(output_path)
        self._ensure_workbook(path)
        if SHEET_B_NAME in self._wb.sheetnames:
            del self._wb[SHEET_B_NAME]
        ws_b = self._wb.create_sheet(SHEET_B_NAME)
        self._write_output_b(ws_b, summary)
        self._atomic_save(path)
        logger.info("Summary sheet written to %s", path)

    def write_excel(
        self,
        output_rows: list[OutputRow],
        scope_summary: ScopeAreaSummary | None,
        output_path: str | Path,
    ) -> None:
        path = Path(output_path)
        self._ensure_workbook(path)
        for row in output_rows:
            self._write_data_row(self._ws_a, row, self._ws_a.max_row + 1)
        if scope_summary:
            if SHEET_B_NAME in self._wb.sheetnames:
                del self._wb[SHEET_B_NAME]
            ws_b = self._wb.create_sheet(SHEET_B_NAME)
            self._write_output_b(ws_b, scope_summary)
        self._atomic_save(path)
        logger.info("Output written to %s (%d rows)", path, len(output_rows))

    def _write_headers(self, ws) -> None:
        for col_idx, (_, label, _w, _wrap) in enumerate(_OUTPUT_A_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER

    def _write_data_row(self, ws, row: OutputRow, row_idx: int) -> None:
        category_val = getattr(row, "risk_category", "") or ""
        color = _category_fill(category_val)
        fill = PatternFill("solid", fgColor=color) if color else None
        for col_idx, (attr, _label, _w, wrap) in enumerate(_OUTPUT_A_COLUMNS, start=1):
            val = _cell_value(row, attr)
            if attr == "source_details" and val and "Source:" in val:
                val = _format_source_details(val)
            if attr in ("risk_description", "risk_rationale") and val:
                val = _apply_bold(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=_xml_safe(val))
            if fill is not None:
                cell.fill = fill
            cell.border = _BORDER
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")
        ws.row_dimensions[row_idx].height = EXCEL_MAX_ROW_HEIGHT

    def write_aggregation_sheets(self, risks: list, scope_overall: dict, output_path) -> None:
        """Write the aggregated risk-analysis sheet (one row per clustered risk) and the grouped
        Risk Summary sheet. `scope_overall` maps scope_area -> overall (averaged) category."""
        path = Path(output_path)
        self._ensure_workbook(path)
        for name in (SHEET_C_NAME, SHEET_D_NAME):
            if name in self._wb.sheetnames:
                del self._wb[name]

        ws_c = self._wb.create_sheet(SHEET_C_NAME)
        self._write_agg_headers(ws_c)
        ws_c.freeze_panes = "A2"
        for col_idx, (_, _, width) in enumerate(_OUTPUT_C_COLUMNS, start=1):
            ws_c.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
        for position, risk in enumerate(risks, start=1):
            risk.overall_risk_category = scope_overall.get(risk.scope_area, "") or ""
            self._write_agg_row(ws_c, risk, position, position + 1)

        ws_d = self._wb.create_sheet(SHEET_D_NAME)
        self._write_risk_summary(ws_d, risks, scope_overall)

        self._atomic_save(path)
        logger.info("Aggregation sheets written to %s (%d risks)", path, len(risks))

    def _write_agg_headers(self, ws) -> None:
        for col_idx, (_, label, _w) in enumerate(_OUTPUT_C_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER

    def _write_agg_row(self, ws, risk, sno: int, row_idx: int) -> None:
        color = _category_fill(getattr(risk, "risk_category", "") or "")
        fill = PatternFill("solid", fgColor=color) if color else None
        for col_idx, (attr, _label, _w) in enumerate(_OUTPUT_C_COLUMNS, start=1):
            val = sno if attr == "sno" else _agg_cell_value(risk, attr)
            if attr == "source_details" and isinstance(val, str) and "Source:" in val:
                val = _format_source_details(val)
            if attr in ("risk_description", "risk_rationale") and isinstance(val, str) and val:
                val = _apply_bold(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=_xml_safe(val))
            if fill is not None:
                cell.fill = fill
            cell.border = _BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = EXCEL_MAX_ROW_HEIGHT

    def _write_risk_summary(self, ws, risks: list, scope_overall: dict) -> None:
        for col_idx, (label, width) in enumerate(_RISK_SUMMARY_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"

        def _rank(risk) -> int:
            return CATEGORY_RANK.get(getattr(risk, "risk_category", ""), 0)

        ordered = sorted(risks, key=lambda r: (r.scope_area, -_rank(r)))
        for row_idx, risk in enumerate(ordered, start=2):
            if getattr(risk, "past_occurrence", False):
                past = "Yes"
            elif not getattr(risk, "past_occurrence_determinable", True):
                past = "Unable to determine based on available evidence"
            else:
                past = "No"
            values = [
                risk.scope_area,
                scope_overall.get(risk.scope_area, risk.risk_category),
                risk.risk_type,
                risk.risk_category,
                past,
                "AI Risk Analysis",
            ]
            color = _category_fill(risk.risk_category)
            fill = PatternFill("solid", fgColor=color) if color else None
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=_xml_safe(value))
                if fill is not None:
                    cell.fill = fill
                cell.border = _BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def write_past_audit_sheet(self, observations: list, output_path) -> None:
        """Write the Past Audit Report sheet — one row per extracted observation."""
        path = Path(output_path)
        self._ensure_workbook(path)
        if SHEET_E_NAME in self._wb.sheetnames:
            del self._wb[SHEET_E_NAME]
        ws = self._wb.create_sheet(SHEET_E_NAME)

        for col_idx, (_, label, width) in enumerate(_OUTPUT_E_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"

        for position, obs in enumerate(observations, start=1):
            for col_idx, (attr, _label, _w) in enumerate(_OUTPUT_E_COLUMNS, start=1):
                val = position if attr == "sno" else (getattr(obs, attr, "") or "")
                if hasattr(val, "strftime"):
                    val = format_date(val)
                elif isinstance(val, str):
                    val = normalize_dates(val)
                cell = ws.cell(row=position + 1, column=col_idx, value=_xml_safe(val))
                cell.border = _BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[position + 1].height = EXCEL_MAX_ROW_HEIGHT

        self._atomic_save(path)
        logger.info("Past Audit Report sheet written to %s (%d observations)", path, len(observations))

    def write_scope_txn_summary(self, rows: list[tuple], output_path) -> None:
        """Write the per-scope transaction summary sheet.

        `rows` = list of 9-tuples:
          (scope_area_display, txns_tables, txns_supp_docs, total_transactions,
           amount_tables, amount_supp_docs, total_amount, table_sources, supp_doc_sources)
        Both transaction counts and amounts are shown separately by origin (tables vs supporting docs)
        plus a combined total; '-' is already rendered for anything with no data."""
        path = Path(output_path)
        self._ensure_workbook(path)
        if SHEET_F_NAME in self._wb.sheetnames:
            del self._wb[SHEET_F_NAME]
        ws = self._wb.create_sheet(SHEET_F_NAME)

        headers = [
            ("Scope Area", 46),
            ("Transactions — Transactional Tables", 18),
            ("Transactions — Supporting Docs", 18),
            ("Total Transactions", 16),
            ("Amount — Transactional Tables (Local / USD)", 40),
            ("Amount — Supporting Docs (Local / USD)", 38),
            ("Total Amount (Local / USD)", 40),
            ("Transactional Table Source(s)", 30),
            ("Supporting-Doc Source(s)", 30),
        ]
        for col_idx, (label, width) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"

        for position, row in enumerate(rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=position, column=col_idx, value=_xml_safe(value))
                cell.border = _BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        self._atomic_save(path)
        logger.info("Scope Txn Summary sheet written to %s (%d scope areas)", path, len(rows))

    def _write_output_b(self, ws, summary: ScopeAreaSummary) -> None:
        rows = [
            ("Scope ID",               summary.scope_id),
            ("Scope Name",             summary.scope_name),
            ("Total Questions",        summary.total_questions),
            ("Potential Risk",         summary.high_risk_count),
            ("Further Review Required", summary.medium_risk_count),
            ("No Risk",                summary.no_risk_count),
            ("Average Risk Category Score", f"{summary.average_risk_score:.2f}"),
            ("Overall Risk Rating",    summary.overall_risk_rating.value),
            ("Transactional Findings", normalize_dates(summary.transactional_findings_summary or "")),
            ("Top Compliance Gaps",    normalize_dates("\n".join(f"• {g}" for g in summary.top_compliance_gaps))),
            ("Generated At",           format_date(summary.generated_at)),
        ]
        for row_idx, (label, value) in enumerate(rows, start=1):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = Font(bold=True)
            label_cell.fill = PatternFill("solid", fgColor="D6E4F0")
            label_cell.border = _BORDER
            val_cell = ws.cell(row=row_idx, column=2, value=_xml_safe(str(value)))
            val_cell.alignment = Alignment(wrap_text=True, vertical="top")
            val_cell.border = _BORDER
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 60
