"""Save the ENGLISH translation of a non-English ingested doc to the translated-docs UC Volume.

The translation is the extracted TEXT (not a byte-for-byte clone of the source). It is rendered in a
format that MATCHES the source family so the stored copy is natural to open:
  * spreadsheets (.xlsx / .xls / .xlsm) -> a real .xlsx workbook ('[Sheet: X]' markers become sheets,
    ' | '-separated rows become cell rows);
  * .csv -> a .csv (each ' | ' row written as comma-separated cells);
  * everything else (prose: pdf / docx / pptx / html / msg / txt) -> a clean .docx.
One file per source doc, per auditor, under:
  <translated_docs_volume_path>/<auditor_id>/<original_stem>__EN.<ext>

Authenticated via the service principal. Best-effort: any failure logs a warning and never breaks the
ingestion run (the translation is still used in-memory regardless).
"""

from __future__ import annotations

import csv
import re
import tempfile
import uuid
from pathlib import Path

from src.config.databricks_auth import build_config
from src.utils.logger import get_logger

logger = get_logger(__name__)

# Source extensions whose translated copy is written as a spreadsheet / csv (not a .docx).
_SPREADSHEET_EXTS = {".xlsx", ".xls", ".xlsm"}
_CSV_EXTS = {".csv"}

# A line emitted by the Excel extractor to mark the start of a sheet: "[Sheet: <name>]".
_SHEET_RE = re.compile(r"^\s*\[Sheet:\s*(.*?)\]\s*$")


def _split_row(line: str) -> list[str]:
    """A ' | '-separated extractor row -> cell list; a plain line -> a single cell."""
    return [c.strip() for c in line.split(" | ")] if " | " in line else [line]


def _parse_sheets(english_text: str) -> list[tuple[str, list[list[str]]]]:
    """Reconstruct (sheet_name, rows) groups from the flattened extractor text. '[Sheet: X]' markers
    start a new sheet; blank lines are dropped. Text with no marker becomes a single 'Sheet1'."""
    sheets: list[tuple[str, list[list[str]]]] = []
    name: str | None = None
    rows: list[list[str]] = []

    def _flush() -> None:
        nonlocal name, rows
        if rows or name is not None:
            sheets.append((name or f"Sheet{len(sheets) + 1}", rows))
        rows = []

    for raw in (english_text or "").split("\n"):
        line = raw.strip()
        marker = _SHEET_RE.match(line)
        if marker:
            _flush()
            name = marker.group(1).strip() or None
            continue
        if not line:
            continue
        rows.append(_split_row(line))
    _flush()
    return sheets or [("Sheet1", [])]


def _safe_sheet_name(name: str, idx: int, used: set[str]) -> str:
    """Excel sheet names: ≤31 chars, none of []:*?/\\, and unique (case-insensitive)."""
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name or "").strip() or f"Sheet{idx}"
    cleaned = cleaned[:31]
    candidate, k = cleaned, 1
    while candidate.lower() in used:
        suffix = f"_{k}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        k += 1
    used.add(candidate.lower())
    return candidate


def _build_xlsx(english_text: str, out_path: Path) -> None:
    """Render the translated text as a real .xlsx (write-only for low memory on large ledgers)."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    wb = Workbook(write_only=True)   # no default sheet; streams rows to disk
    used: set[str] = set()
    for idx, (name, rows) in enumerate(_parse_sheets(english_text), start=1):
        ws = wb.create_sheet(_safe_sheet_name(name, idx, used))
        for row in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", c) for c in row])
    wb.save(str(out_path))


def _build_csv(english_text: str, out_path: Path) -> None:
    """Render the translated text as a .csv (utf-8-sig so Excel opens non-ASCII cleanly)."""
    with open(out_path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        for raw in (english_text or "").split("\n"):
            line = raw.strip()
            if not line or _SHEET_RE.match(line):
                continue
            writer.writerow(_split_row(line))


def _build_docx(english_text: str, out_path: Path) -> None:
    """Render the translated text as a .docx: ' | ' rows become Word tables, everything else prose."""
    from docx import Document

    doc = Document()
    pending: list[list[str]] = []

    def _flush_table() -> None:
        if not pending:
            return
        cols = max(len(r) for r in pending)
        table = doc.add_table(rows=0, cols=cols)
        try:
            table.style = "Table Grid"
        except Exception:  # noqa: BLE001 — style may not exist in the default template
            pass
        for row in pending:
            cells = table.add_row().cells
            for j, value in enumerate(row):
                cells[j].text = value
        doc.add_paragraph("")
        pending.clear()

    for line in (english_text or "").split("\n"):
        stripped = line.strip()
        if " | " in stripped:
            pending.append([c.strip() for c in stripped.split(" | ")])
            continue
        _flush_table()
        doc.add_paragraph(stripped)
    _flush_table()
    doc.save(str(out_path))


def _renderer_for(original_name: str):
    """Pick (output extension, builder) from the SOURCE file's extension."""
    suffix = Path(original_name).suffix.lower()
    if suffix in _SPREADSHEET_EXTS:
        return ".xlsx", _build_xlsx
    if suffix in _CSV_EXTS:
        return ".csv", _build_csv
    return ".docx", _build_docx


def save_translated_doc(
    settings, original_name: str, english_text: str, auditor_id: str | None = None
) -> str | None:
    """Build a translated copy of `english_text` in the format matching the source (xlsx/csv/docx) and
    upload it to the translated-docs Volume. Returns the dest path or None. Never raises."""
    base = (getattr(settings, "translated_docs_volume_path", "") or "").rstrip("/")
    if not base:
        return None
    if not (english_text or "").strip():
        return None
    auditor = (auditor_id or getattr(settings, "auditor_id", "") or "").strip()
    stem = Path(original_name).stem
    ext, builder = _renderer_for(original_name)
    dest = f"{base}/{auditor}/{stem}__EN{ext}" if auditor else f"{base}/{stem}__EN{ext}"

    tmp = Path(tempfile.gettempdir()) / f"translated_{uuid.uuid4().hex}{ext}"
    try:
        builder(english_text, tmp)
        from databricks.sdk import WorkspaceClient

        client = WorkspaceClient(config=build_config(settings))
        with open(tmp, "rb") as fh:
            client.files.upload(dest, fh, overwrite=True)
        logger.info("[TranslatedDocs] Uploaded %s__EN%s -> %s", stem, ext, dest)
        return dest
    except Exception as exc:  # noqa: BLE001 — best-effort; translation is still used in-memory
        logger.warning("[TranslatedDocs] Save failed (translation still used) | dest=%s | %s", dest, exc)
        return None
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except Exception:  # noqa: BLE001
            pass


# Backward-compatible alias (the function used to always emit a .docx).
save_translated_docx = save_translated_doc
