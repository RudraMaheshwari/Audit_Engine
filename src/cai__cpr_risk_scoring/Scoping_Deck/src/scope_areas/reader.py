"""Reads the AI Scope (risk-scope) workbook produced by the AI Scope pipeline.

The sheet carries one row per scope area with headers: Scope Area | Scope Decision |
Overall Risk Category | Proposed Sample Size | Summary of Risks. Columns are matched by HEADER NAME
(not position), so added/re-ordered columns don't break the read. Returns [] if the file is absent.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from src.utils.logger import get_logger

logger = get_logger(__name__)


@dataclass(frozen=True)
class ScopeAreaRow:
    scope_area: str
    scope_decision: str          # "Included" | "Excluded" (default Excluded)
    overall_risk_category: str
    proposed_sample_size: str
    summary_of_risks: str

    @property
    def is_included(self) -> bool:
        return self.scope_decision.strip().lower() == "included"


# Header aliases (lowercased) → canonical field. Tolerant of minor wording differences.
_HEADER_MAP = {
    "scope area": "scope_area",
    "compliance process review area": "scope_area",
    "scope decision": "scope_decision",
    "scope status": "scope_decision",
    "overall risk category": "overall_risk_category",
    "proposed sample size": "proposed_sample_size",
    "sample size": "proposed_sample_size",
    "summary of risks": "summary_of_risks",
    "comments": "summary_of_risks",
}


_DB_SOURCES = {"db_table", "db", "delta", "databricks", "table"}


def load_scope_areas(path: str, sheet: str, *, source: str = "excel", settings=None) -> list[ScopeAreaRow]:
    """Load scope-area rows from the AI Scope output — an Excel workbook (source='excel') or a Delta
    table (source='db_table'). `path` is the workbook path or the table FQN accordingly. Returns []
    if the input is missing/unreadable."""
    if not path:
        logger.info("[ScopeAreas] No risk-scope path configured — slides 9/10 will be skipped")
        return []

    if (source or "excel").strip().lower() in _DB_SOURCES:
        header, data = _read_table(path, settings)
    else:
        header, data = _read_excel(path, sheet)
    if header is None:
        return []
    return _build_rows(header, data, path)


def _read_excel(path: str, sheet: str):
    p = Path(path)
    if not p.is_file():
        logger.info("[ScopeAreas] Risk-scope file not found (%s) — slides 9/10 will be skipped", path)
        return None, []
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("[ScopeAreas] Could not open risk-scope workbook %s | %s", path, exc)
        return None, []
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        logger.warning("[ScopeAreas] Could not read sheet %r | %s", sheet, exc)
        return None, []
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return (rows[0], rows[1:]) if rows else (None, [])


def _read_table(table: str, settings):
    """Read the AI Scope rows from a Delta table (the `_copy` output table)."""
    if settings is None:
        logger.warning("[ScopeAreas] db_table source but no settings — cannot read %s", table)
        return None, []
    try:
        from sqlalchemy import text

        from src.config.databricks_auth import build_engine

        engine = build_engine(settings)
        with engine.connect() as conn:
            result = conn.execute(text(f"SELECT * FROM {table}"))
            header = [str(c) for c in result.keys()]
            data = [tuple(r) for r in result.fetchall()]
        return header, data
    except Exception as exc:
        logger.warning("[ScopeAreas] Could not read risk-scope table %s | %s", table, exc)
        return None, []


def _build_rows(header, data_rows, src_label: str) -> list[ScopeAreaRow]:
    col_index: dict[str, int] = {}
    for idx, label in enumerate(header):
        # Normalise 'Scope_Area' / 'scope area' alike so Delta (underscored) + Excel headers both map.
        key = str(label or "").strip().lower().replace("_", " ")
        field = _HEADER_MAP.get(key)
        if field and field not in col_index:
            col_index[field] = idx

    if "scope_area" not in col_index:
        logger.warning("[ScopeAreas] No 'Scope Area' column in %s — cannot build slides 9/10", src_label)
        return []

    def cell(row, field: str) -> str:
        idx = col_index.get(field)
        if idx is None or idx >= len(row):
            return ""
        val = row[idx]
        return "" if val is None else str(val).strip()

    out: list[ScopeAreaRow] = []
    for row in data_rows:
        scope_area = cell(row, "scope_area")
        if not scope_area:
            continue
        out.append(ScopeAreaRow(
            scope_area=scope_area,
            scope_decision=cell(row, "scope_decision") or "Excluded",
            overall_risk_category=cell(row, "overall_risk_category"),
            proposed_sample_size=cell(row, "proposed_sample_size"),
            summary_of_risks=cell(row, "summary_of_risks"),
        ))
    logger.info("[ScopeAreas] Loaded %d scope-area row(s) from %s", len(out), src_label)
    return out
