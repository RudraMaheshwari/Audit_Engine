"""Load per-scope transaction totals from Output F (Scope Txn Summary) for slides 9/10.

Output F lives alongside Output E — a sheet in the same audit workbook (excel source) or a sibling Delta
table (db_table source). We read the per-scope Total Transactions and the full Total Amount cell (both
local currency and USD, verbatim), keyed by a normalised scope name so the deck's scope areas resolve
regardless of exact spelling. Best-effort: returns {} on any failure.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from sqlalchemy import text

from src.utils.logger import get_logger

from ..config import ScopingDeckConfig

logger = get_logger(__name__)

_MISSING = "-"
_SOURCE_EXCEL = "excel"

# Normalised header -> internal field (resolves both Excel display headers and snake_case DB columns).
_FIELD_BY_KEY = {
    "scopearea": "scope_area",
    "totaltransactions": "total_txns",
    "totalamountlocalusd": "total_amount",
}


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def _amount_text(amount_cell) -> str:
    """Return the Total Amount cell verbatim (both local currency and USD); '-' when blank."""
    s = str(amount_cell or "").strip()
    return s if s else _MISSING


def _resolve_columns(headers: list) -> dict[int, str]:
    return {
        idx: _FIELD_BY_KEY[_normalize(str(h))]
        for idx, h in enumerate(headers)
        if _normalize(str(h)) in _FIELD_BY_KEY
    }


def _records_from_excel(config: ScopingDeckConfig) -> list[dict]:
    path = Path((config.output_f_path or "").strip())
    if not path.exists():
        logger.warning("[OutputF] Workbook not found: %s", path)
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        logger.warning("[OutputF] Could not open workbook %s | %s", path, exc)
        return []
    try:
        target = _normalize(config.output_f_sheet)
        ws = next((wb[s] for s in wb.sheetnames if _normalize(s) == target), None)
        if ws is None:
            logger.info("[OutputF] No '%s' sheet in workbook — scope txn totals unavailable", config.output_f_sheet)
            return []
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    col_field = _resolve_columns([str(h or "") for h in rows[0]])
    if "scope_area" not in col_field.values():
        logger.warning("[OutputF] Output F sheet missing a Scope Area column — headers=%s", list(rows[0]))
        return []
    return [
        {f: (row[idx] if idx < len(row) else None) for idx, f in col_field.items()}
        for row in rows[1:]
    ]


def _records_from_db(config: ScopingDeckConfig) -> list[dict]:
    from src.config.databricks_auth import build_engine

    table = (config.output_f_path or "").strip()
    engine = build_engine(config.settings, schema=config.settings.databricks_schema_name)
    with engine.connect() as conn:
        result = conn.execute(text(f"SELECT * FROM {table}"))
        col_field = _resolve_columns([str(c) for c in result.keys()])
        fetched = result.fetchall()
    if "scope_area" not in col_field.values():
        logger.warning("[OutputF] Output F table %r missing a Scope Area column", table)
        return []
    return [{f: row[idx] for idx, f in col_field.items()} for row in fetched]


def load_output_f_totals(config: ScopingDeckConfig) -> dict[str, tuple[str, str]]:
    """Return {normalized_scope_area: (total_transactions, amount)} from Output F — read from the same
    workbook as Output E (excel source) or the Output F Delta table (db_table source). {} on failure."""
    path_str = (getattr(config, "output_f_path", "") or "").strip()
    if not path_str:
        logger.info("[OutputF] No Output F path configured — scope txn totals render as '-'")
        return {}
    try:
        # Output F follows the same source as Output E (they are co-located).
        if config.past_review_source.strip().lower() == _SOURCE_EXCEL:
            records = _records_from_excel(config)
        else:
            records = _records_from_db(config)
    except Exception as exc:  # noqa: BLE001 — scope txn totals are best-effort
        logger.warning("[OutputF] Scope txn totals unavailable | %s", exc)
        return {}

    out: dict[str, tuple[str, str]] = {}
    for rec in records:
        scope = str(rec.get("scope_area") or "").strip()
        if not scope:
            continue
        count = str(rec.get("total_txns") or "").strip() or _MISSING
        amount = _amount_text(rec.get("total_amount"))
        out[_normalize(scope)] = (count, amount)

    logger.info("[OutputF] Loaded transaction totals for %d scope area(s)", len(out))
    return out
