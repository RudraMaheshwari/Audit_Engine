"""Fetches CPI figures from Transparency International, ephemerally and without a browser.

The country page's HTML links the current data file (whatever TI named it), so the file URL is
discovered rather than guessed. The file is downloaded into memory, its non-standard OOXML namespace
is normalised so openpyxl can read it, the two most recent years are read from the timeseries sheet,
and the bytes are discarded — nothing is written to disk.
"""

from __future__ import annotations

import io
import os
import re
import warnings
import zipfile

import requests

from src.utils.logger import get_logger

from .models import CpiResult, CpiYearFigure

logger = get_logger(__name__)

warnings.filterwarnings("ignore", module="openpyxl")

# Corporate networks often intercept HTTPS with an internal CA that Python's default bundle does not
# trust. truststore makes requests use the OS trust store (which holds the corporate CA), fixing the
# "unable to get local issuer certificate" error with no config. Best-effort: ignored if unavailable.
try:
    import truststore

    truststore.inject_into_ssl()
    logger.info("[CPI] Using OS trust store (truststore) for TLS verification")
except Exception:
    pass


def _resolve_verify():
    """TLS verification setting: a CA bundle path, False (insecure), or True (default)."""
    ca_bundle = (os.getenv("SCOPING_DECK_CA_BUNDLE", "") or "").strip()
    if ca_bundle:
        return ca_bundle
    if (os.getenv("SCOPING_DECK_CPI_INSECURE", "") or "").strip().lower() in ("1", "true", "yes"):
        try:
            from urllib3.exceptions import InsecureRequestWarning

            warnings.simplefilter("ignore", InsecureRequestWarning)
        except Exception:
            pass
        return False
    return True


class CpiRepository:
    """Resolves, downloads, and parses the CPI dataset for a single country."""

    _PAGE_TEMPLATE = "https://www.transparency.org/en/cpi/{year}"
    _USER_AGENT = {"User-Agent": "Mozilla/5.0"}
    _XLSX_PATTERN = re.compile(r"https?://[^\"')\s]+?\.xlsx")
    _TIMESERIES_HINT = "timeseries"
    _NAMESPACE_FIXES = {
        b"http://purl.oclc.org/ooxml/spreadsheetml/main":
            b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        b"http://purl.oclc.org/ooxml/officeDocument/relationships":
            b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        b"http://purl.oclc.org/ooxml/package/relationships":
            b"http://schemas.openxmlformats.org/package/2006/relationships",
    }

    def __init__(self, timeout_seconds: int = 40) -> None:
        self._timeout = timeout_seconds
        self._verify = _resolve_verify()

    def fetch(self, iso3: str, country: str, audit_year: int) -> CpiResult | None:
        """Return the two most recent CPI years for the country, or None if unavailable.

        CPI for year Y is published the following year, so the latest file is for audit_year - 1
        (with a fallback to audit_year - 2 if that year is not yet published).
        """
        for target_year in (audit_year - 1, audit_year - 2):
            url = self._resolve_file_url(target_year)
            if not url:
                continue
            try:
                workbook = self._download_workbook(url)
                result = self._extract(workbook, iso3, country)
            except Exception as exc:
                logger.warning("[CPI] Failed to read %s | %s", url, exc)
                continue
            if result:
                return result
        logger.warning("[CPI] No CPI data resolved for %s (audit_year=%d)", country, audit_year)
        return None

    def _resolve_file_url(self, year: int) -> str | None:
        page = self._PAGE_TEMPLATE.format(year=year)
        try:
            html = requests.get(
                page, headers=self._USER_AGENT, timeout=self._timeout, verify=self._verify
            ).text
        except Exception as exc:
            logger.warning("[CPI] Could not load CPI page %s | %s", page, exc)
            return None
        match = self._XLSX_PATTERN.search(html)
        if not match:
            logger.info("[CPI] No data file linked on %s (year not published yet?)", page)
            return None
        logger.info("[CPI] Resolved data file for %d -> %s", year, match.group())
        return match.group()

    def _download_workbook(self, url: str):
        import openpyxl

        raw = requests.get(
            url, headers=self._USER_AGENT, timeout=self._timeout, verify=self._verify
        ).content
        if len(raw) < 10000:
            raise ValueError(f"downloaded file too small to be a workbook ({len(raw)} bytes)")

        source = zipfile.ZipFile(io.BytesIO(raw))
        fixed = io.BytesIO()
        with zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.endswith((".xml", ".rels")):
                    for old, new in self._NAMESPACE_FIXES.items():
                        data = data.replace(old, new)
                target.writestr(item, data)
        fixed.seek(0)
        return openpyxl.load_workbook(fixed, data_only=True)

    def _extract(self, workbook, iso3: str, country: str) -> CpiResult | None:
        sheet = self._timeseries_sheet(workbook)
        header_row, iso3_col, country_col, year_columns = self._map_columns(sheet)
        if not year_columns:
            return None

        recent_years = sorted(year_columns, reverse=True)[:2]
        totals = {y: self._count_ranked(sheet, header_row, year_columns[y]["rank"]) for y in recent_years}

        row = self._find_country_row(sheet, header_row, iso3_col, country_col, iso3, country)
        if row is None:
            return None

        figures: list[CpiYearFigure] = []
        for year in recent_years:
            score = sheet.cell(row=row, column=year_columns[year]["score"]).value
            rank = sheet.cell(row=row, column=year_columns[year]["rank"]).value
            if not isinstance(score, (int, float)) or not isinstance(rank, (int, float)):
                continue
            figures.append(CpiYearFigure(
                year=year, score=int(score), rank=int(rank), total_countries=totals[year],
            ))
        if not figures:
            return None
        return CpiResult(country=country, figures=tuple(figures))

    def _timeseries_sheet(self, workbook):
        for name in workbook.sheetnames:
            if self._TIMESERIES_HINT in name.lower():
                return workbook[name]
        return workbook[workbook.sheetnames[0]]

    @staticmethod
    def _map_columns(sheet):
        score_re = re.compile(r"CPI score\s*(\d{4})", re.IGNORECASE)
        rank_re = re.compile(r"Rank\s*(\d{4})", re.IGNORECASE)
        for header_row in range(1, 9):
            iso3_col = country_col = None
            year_columns: dict[int, dict[str, int]] = {}
            for col in range(1, sheet.max_column + 1):
                text = str(sheet.cell(row=header_row, column=col).value or "").strip()
                lowered = text.lower()
                if lowered == "iso3":
                    iso3_col = col
                elif "country" in lowered:
                    country_col = col
                score_match = score_re.search(text)
                rank_match = rank_re.search(text)
                if score_match:
                    year_columns.setdefault(int(score_match.group(1)), {})["score"] = col
                if rank_match:
                    year_columns.setdefault(int(rank_match.group(1)), {})["rank"] = col
            complete = {y: c for y, c in year_columns.items() if "score" in c and "rank" in c}
            if iso3_col and complete:
                return header_row, iso3_col, country_col, complete
        return 0, None, None, {}

    @staticmethod
    def _count_ranked(sheet, header_row: int, rank_col: int) -> int:
        return sum(
            1 for r in range(header_row + 1, sheet.max_row + 1)
            if isinstance(sheet.cell(row=r, column=rank_col).value, (int, float))
        )

    # Common country-name aliases → the canonical form, so an input like "Korea" matches the CPI
    # dataset's "South Korea" / "Korea, South" / "Korea, Rep." row (Korea is the known miss; others
    # are added defensively). North Korea is handled separately in _canonical_country.
    _COUNTRY_ALIASES = {
        "republic of korea": "south korea",
        "korea, republic of": "south korea",
        "korea rep": "south korea",
        "korea (south)": "south korea",
        "usa": "united states",
        "uk": "united kingdom",
        "uae": "united arab emirates",
    }

    @classmethod
    def _canonical_country(cls, name: str) -> str:
        """Normalise a country name so equivalent spellings match (e.g. 'Korea' == 'South Korea')."""
        n = " ".join(str(name or "").strip().lower().replace(".", "").replace(",", " ").split())
        if "korea" in n and not any(w in n for w in ("north", "dprk", "democratic")):
            return "south korea"   # 'Korea', 'Korea, South', 'Korea Rep', 'Republic of Korea' → one form
        return cls._COUNTRY_ALIASES.get(n, n)

    @classmethod
    def _find_country_row(cls, sheet, header_row, iso3_col, country_col, iso3, country) -> int | None:
        iso3_target = iso3.strip().upper()
        country_target = cls._canonical_country(country)
        for r in range(header_row + 1, sheet.max_row + 1):
            code = str(sheet.cell(row=r, column=iso3_col).value or "").strip().upper()
            if code == iso3_target:
                return r
            if country_col:
                name = str(sheet.cell(row=r, column=country_col).value or "")
                if cls._canonical_country(name) == country_target:
                    return r
        return None
