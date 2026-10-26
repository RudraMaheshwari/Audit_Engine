"""Loads Output E (Past Audit Report / 'Summary of Past Reviews') from Excel or a Delta table.

Mirrors the AI Scope loader: headers are matched by a normalised key so both the Excel display
headers and the snake_case DB columns resolve the same.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from sqlalchemy import text
from src.config.settings import Settings
from src.utils.logger import get_logger

from ..config import ScopingDeckConfig
from .classifier import PastReviewClassifier
from .models import PastReviewObservation

logger = get_logger(__name__)

_SOURCE_EXCEL = "excel"

_FIELD_BY_KEY = {
    "scopearea": "scope_area",
    "issuedescription": "issue_description",
    "criticality": "criticality",
    "issuetype": "issue_type",
    "sourcedocument": "source_document",
    "auditreport": "audit_report",
    "reportdate": "report_date",
}


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def _year_from(value) -> int | None:
    match = re.search(r"(19|20)\d{2}", str(value or ""))
    return int(match.group()) if match else None


def _to_observation(record: dict) -> PastReviewObservation:
    source_document = str(record.get("source_document") or "").strip()
    audit_report = str(record.get("audit_report") or "").strip()
    kind = PastReviewClassifier.classify(source_document, audit_report)
    return PastReviewObservation(
        scope_area=str(record.get("scope_area") or "").strip(),
        issue_description=str(record.get("issue_description") or "").strip(),
        criticality=str(record.get("criticality") or "").strip(),
        issue_type=str(record.get("issue_type") or "").strip(),
        source_document=source_document,
        audit_report=audit_report,
        report_year=_year_from(record.get("report_date")),
        review_kind=kind,
    )


class OutputELoader:
    """Reads Output E rows into classified PastReviewObservation objects."""

    def __init__(self, config: ScopingDeckConfig) -> None:
        self._config = config

    def load(self) -> list[PastReviewObservation]:
        if self._config.past_review_source.strip().lower() == _SOURCE_EXCEL:
            records = self._load_excel()
        else:
            records = self._load_db_table()
        observations = [_to_observation(r) for r in records if str(r.get("scope_area") or "").strip()]
        logger.info("[PastReview] Loaded %d observation(s) from Output E", len(observations))
        return observations

    def _resolve_columns(self, headers: list[str]) -> dict[int, str]:
        mapping: dict[int, str] = {}
        for idx, header in enumerate(headers):
            field = _FIELD_BY_KEY.get(_normalize(str(header)))
            if field:
                mapping[idx] = field
        return mapping

    def _load_excel(self) -> list[dict]:
        path = Path(self._config.past_review_path)
        if not path.exists():
            raise FileNotFoundError(f"Output E workbook not found: {path}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = self._config.past_review_excel_sheet
            ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[-1]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        if not rows:
            return []
        col_field = self._resolve_columns([str(h or "") for h in rows[0]])
        return [
            {field: row[idx] if idx < len(row) else None for idx, field in col_field.items()}
            for row in rows[1:]
        ]

    def _load_db_table(self) -> list[dict]:
        engine = self._build_engine(self._config.settings)
        with engine.connect() as conn:
            result = conn.execute(text(f"SELECT * FROM {self._config.past_review_path}"))
            col_field = self._resolve_columns([str(c) for c in result.keys()])
            fetched = result.fetchall()
        return [{field: row[idx] for idx, field in col_field.items()} for row in fetched]

    @staticmethod
    def _build_engine(settings: Settings):
        from src.config.databricks_auth import build_engine
        return build_engine(settings, schema=settings.databricks_schema_name)
