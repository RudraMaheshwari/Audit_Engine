"""
Re-run ONLY Phase 3 (transactional analysis) against an existing output workbook.

Phase 3 now produces ONE consolidated row per scope area (~48 rows): each scope is analysed
against the transactional data using its CPR Data Calculations (or the AI's own FCPA intelligence
where the guidance is thin), and — when a risk is found — samples sized to the scope's Expected
Sample Size + buffer are attached in the Samples column.

This script does NOT touch the Phase 1&2 questionnaire rows. It:
  1. Runs TransactionalDataAnalysisAgent.run_all() to produce the fresh per-scope rows
  2. Deletes the OLD Phase 3 rows from "Output A" (identified by the transactional marker in the
     Source Details column) — so re-runs don't stack up duplicates
  3. Appends the fresh rows (renumbered to continue after the questionnaire rows)
  4. Rebuilds the "Output B" summary sheet from the corrected data

Note: run_all() analyses the scopes concurrently (a thread pool of _SCOPE_ANALYSIS_WORKERS),
so a full re-run is much faster than one-at-a-time — but still expect a few minutes for ~48 scopes.

Usage:
    python scripts/rerun_phase3.py
    python scripts/rerun_phase3.py --file FCPA_Audit_Sam_1_20260903_100017.xlsx
    python scripts/rerun_phase3.py --file C:\\full\\path\\to\\output.xlsx

If --file is a bare name it is resolved against the project's output/ directory.
"""
import argparse
import sys
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

def _find_project_root(start: Path) -> Path:
    """Walk up from this file until we find the folder that contains 'src'.
    Works whether the script lives in scripts/ or directly at the project root."""
    for p in [start, *start.parents]:
        if (p / "src").is_dir():
            return p
    return start


_PROJECT_ROOT = _find_project_root(Path(__file__).resolve().parent)
sys.path.insert(0, str(_PROJECT_ROOT))

import openpyxl

from src.config.constants import SHEET_A_NAME
from src.config.settings import get_settings
from src.agents.transactional_data_analysis_agent import TransactionalDataAnalysisAgent
from src.infrastructure.llm.azure_openai_provider import AzureOpenAILLMProvider
from src.output.output_writer import OutputWriter
from src.repositories.output_repo import OutputRepository
from src.utils.logger import get_logger

logger = get_logger(__name__)

_DEFAULT_FILE = "FCPA_Audit_Sam_1_20260903_100017.xlsx"
_SCOPE_NAME   = "TPI Due Diligence"
# Phase 3 rows are tagged with this in their Source Details column (see _to_risk_output).
_PHASE3_MARKER = "Transactional Data Analysis"


def _resolve_path(file_arg: str) -> Path:
    p = Path(file_arg)
    if p.is_absolute() or p.exists():
        return p
    return _PROJECT_ROOT / "output" / file_arg


def _delete_old_phase3_rows(path: Path, keep_rows: int | None = None) -> int:
    """Delete the entire previous Phase 3 block from Output A (the tail after the questionnaire
    rows) and return how many rows were removed.

    Boundary detection:
      - If --keep-rows N is given, keep the first N data rows (header + N) and delete everything
        after them (i.e. delete rows N+2 .. end).
      - Otherwise, find the FIRST row tagged with the transactional marker in Source Details and
        delete from there to the end. Questionnaire rows never carry that marker, so they are safe.
    """
    wb = openpyxl.load_workbook(path)
    if SHEET_A_NAME not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"Sheet '{SHEET_A_NAME}' not found in {path.name}")
    ws = wb[SHEET_A_NAME]

    header = [str(c.value or "").strip().lower() for c in ws[1]]
    try:
        src_col = header.index("source details") + 1  # 1-based
    except ValueError:
        wb.close()
        raise RuntimeError("Could not find a 'Source Details' column to locate Phase 3 rows")

    if keep_rows is not None and keep_rows > 0:
        first_p3 = keep_rows + 2  # row 1 = header, rows 2..keep_rows+1 = questionnaire
    else:
        first_p3 = None
        for r in range(2, ws.max_row + 1):
            if _PHASE3_MARKER in str(ws.cell(row=r, column=src_col).value or ""):
                first_p3 = r
                break

    if not first_p3 or first_p3 > ws.max_row:
        wb.close()
        logger.info("[Rerun P3] No previous Phase 3 rows found to delete.")
        return 0

    count = ws.max_row - first_p3 + 1
    ws.delete_rows(first_p3, count)
    wb.save(path)
    wb.close()
    logger.info("[Rerun P3] Deleted %d old Phase 3 row(s) from %s (rows %d..end)",
                count, SHEET_A_NAME, first_p3)
    return count


def main() -> None:
    parser = argparse.ArgumentParser(description="Re-run only Phase 3 on an existing output file")
    parser.add_argument("--file", default=_DEFAULT_FILE,
                        help="Output workbook (bare name resolves under output/, or a full path)")
    parser.add_argument("--keep-rows", type=int, default=None,
                        help="Keep the first N data rows (the questionnaire rows, e.g. 428) and "
                             "delete everything after them. If omitted, the Phase 3 block is "
                             "detected by its transactional marker.")
    args = parser.parse_args()

    path = _resolve_path(args.file)
    if not path.exists():
        raise FileNotFoundError(f"Output file not found: {path}")
    logger.info("[Rerun P3] Target workbook: %s", path)

    # 1. Fresh Phase 3 rows — one consolidated assessment per scope area.
    settings = get_settings()
    llm = AzureOpenAILLMProvider(settings).get_client()
    agent = TransactionalDataAnalysisAgent(llm, settings)
    logger.info("[Rerun P3] Running Phase 3 (scopes analysed concurrently) ...")
    findings = agent.run_all()
    logger.info("[Rerun P3] Phase 3 produced %d per-scope row(s)", len(findings))

    # 2. Remove the previous Phase 3 block so we don't duplicate.
    _delete_old_phase3_rows(path, keep_rows=args.keep_rows)

    # 3 & 4. Append the fresh rows and rebuild the summary. With the file already present,
    # OutputWriter.write() runs in resume mode: it keeps the questionnaire rows, numbers the
    # new analysis rows after them, and rebuilds Output B from the file.
    writer = OutputWriter(OutputRepository())
    writer.write(
        states=[],
        analysis_findings=findings,
        scope_name=_SCOPE_NAME,
        output_path=path,
    )
    logger.info("[Rerun P3] DONE — appended %d per-scope row(s) and rebuilt the summary | file=%s",
                len(findings), path.name)


if __name__ == "__main__":
    main()
