"""Update currency conversion in Output F (Scope Txn Summary) of an existing audit workbook.

Reads the existing Scope Txn Summary sheet, finds rows with a plain "USD X,XXX.XX" amount
(not yet converted), fetches the FX rate for Dec 31 of the audit year, and rewrites the amount
cell as "<LOCAL> X,XXX.XX / USD X,XXX.XX".  All other sheets are untouched.

Usage (from the project root):
    python -m scripts.update_output_f --output <path_to_xlsx> --currency AUD
    python -m scripts.update_output_f --output <path_to_xlsx> --currency AUD --year 2024
"""

from __future__ import annotations

import argparse
import datetime as _dt
import os
import re
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from openpyxl.styles import Alignment, Border, Side

from src.services.fx_rates import FxRateService
from src.utils.logger import get_logger

logger = get_logger(__name__)

_SHEET_F_NAME = "Output F - Scope Txn Summary"
_USD_RE = re.compile(r"^USD\s*([\d,]+(?:\.\d+)?)$", re.IGNORECASE)
_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Update Output F currency conversion in-place")
    parser.add_argument("--output",   required=True, help="Path to the existing audit output .xlsx")
    parser.add_argument("--currency", required=True, help="Local currency code, e.g. AUD")
    parser.add_argument("--year",     type=int, default=None,
                        help="Audit year (default: current year - 1). Rate used: Dec 31 of that year.")
    return parser.parse_args()


def main() -> None:
    args      = _parse_args()
    path      = args.output
    local     = args.currency.upper()
    audit_year = args.year or (_dt.date.today().year - 1)
    rate_date  = _dt.date(audit_year, 12, 31)

    if not os.path.isfile(path):
        print(f"ERROR: File not found: {path}")
        sys.exit(1)

    fx = FxRateService(base="USD", local=local)
    fx.load_period(rate_date, rate_date)
    rate = fx.rate_on(rate_date)
    if rate is None:
        print(f"ERROR: Could not load USD->{local} rate for {rate_date}. Check network/SSL.")
        sys.exit(1)

    logger.info("[UpdateOutputF] USD->%s rate on %s = %.4f", local, rate_date, rate)

    wb = openpyxl.load_workbook(path)
    if _SHEET_F_NAME not in wb.sheetnames:
        print(f"ERROR: Sheet '{_SHEET_F_NAME}' not found in {path}")
        sys.exit(1)

    ws = wb[_SHEET_F_NAME]
    updated = 0
    for row in ws.iter_rows(min_row=2):
        amount_cell = row[2]
        raw = str(amount_cell.value or "").strip()
        m = _USD_RE.match(raw)
        if not m:
            continue
        usd_amount = float(m.group(1).replace(",", ""))
        local_amount = usd_amount * rate
        amount_cell.value = f"{local} {local_amount:,.2f} / USD {usd_amount:,.2f}"
        amount_cell.alignment = Alignment(wrap_text=True, vertical="top")
        amount_cell.border = _BORDER
        updated += 1

    wb.save(path)
    logger.info("[UpdateOutputF] Done | updated=%d rows | rate=%.4f | file=%s", updated, rate, path)
    print(f"\nUpdated {updated} row(s) — USD->{local} @ {rate:.4f} ({rate_date}) -> {path}")


if __name__ == "__main__":
    main()
